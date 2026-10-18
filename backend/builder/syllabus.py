"""A syllabus skeleton (.docx) for a build, made with python-docx.

The document fills in what Aletheia knows (title, term, outcomes, the dated schedule, major
assignments, credits) and leaves highlighted, bracketed placeholders for everything the
instructor writes. Its own text stays short and neutral.
"""
from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from builder import dated
from catalog.models import OutcomeSet

POLICIES = [
    ("Attendance", "[Attendance policy]"),
    ("Late work", "[Late work policy]"),
    ("Academic integrity", "[Academic integrity policy]"),
    ("Generative AI", "[Policy on generative AI tools]"),
    ("Accessibility", "[Accessibility statement and disability services contact]"),
    ("Title IX and basic needs", "[Title IX, basic needs, and other required campus statements]"),
]
MAJOR_ROLES = ("major", "short_term")


def _placeholder(paragraph, text: str):
    run = paragraph.add_run(text)
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return run


def _field(doc, label: str, value: str = "", placeholder: str = ""):
    p = doc.add_paragraph()
    p.add_run(f"{label}: ").bold = True
    if value:
        p.add_run(value)
    else:
        _placeholder(p, placeholder)
    return p


def _hyperlink(paragraph, url: str, text: str | None = None) -> None:
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(color)
    props.append(underline)
    run.append(props)
    t = OxmlElement("w:t")
    t.text = text or url
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def _numbered(doc, texts: list[str]) -> list:
    """A numbered list that starts at 1. Every "List Number" paragraph in python-docx's template
    shares one numbering instance, so a second list would carry on from the first without this."""
    numbering = doc.part.numbering_part.element
    style_num = doc.styles["List Number"].element.pPr.numPr.numId.val
    abstract = numbering.num_having_numId(style_num).abstractNumId.val
    num = numbering.add_num(abstract)
    num.add_lvlOverride(ilvl=0).add_startOverride(1)
    out = []
    for t in texts:
        p = doc.add_paragraph(t, style="List Number")
        num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
        num_pr.get_or_add_ilvl().val = 0
        num_pr.get_or_add_numId().val = num.numId
        out.append(p)
    return out


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def _cell_lines(cell, lines: list[str], placeholder: str = "") -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    if not lines:
        if placeholder:
            _placeholder(p, placeholder)
        return
    for i, line in enumerate(lines):
        if i:
            p = cell.add_paragraph()
        p.add_run(line)


def outcome_sets(state: dict) -> list[tuple[OutcomeSet, list[str]]]:
    """Outcome sets chosen for the build (``outcome_set_ids``), in that order, with their texts."""
    ids = [i for i in (state.get("outcome_set_ids") or []) if isinstance(i, str)]
    if not ids:
        return []
    found = {o.id: o for o in OutcomeSet.objects.filter(id__in=ids).prefetch_related("outcomes")}
    out = []
    for i in ids:
        o = found.get(i)
        if o:
            texts = [x.text for x in o.outcomes.all()]
            if texts:
                out.append((o, texts))
    return out


def _schedule_table(doc, ctx: dict) -> None:
    rows = ctx.get("schedule_rows") or dated.schedule_rows(ctx, None)
    has_dates = bool(ctx.get("schedule"))
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    widths = [Inches(0.6), Inches(1.2), Inches(1.5), Inches(2.3), Inches(1.4)]
    header = table.rows[0]
    for cell, text, width in zip(header.cells, ["Week", "Dates", "Unit", "Readings", "Assignments due"], widths):
        cell.text = ""
        cell.paragraphs[0].add_run(text).bold = True
        cell.width = width
    _repeat_header(header)
    for row in rows:
        cells = table.add_row().cells
        for cell, width in zip(cells, widths):
            cell.width = width
        if row["kind"] == "break":
            cells[0].text = ""
            cells[1].text = row["dates"]
            merged = cells[2].merge(cells[4])
            merged.text = ""
            merged.paragraphs[0].add_run(f"No class: {row['unit']}").italic = True
            continue
        cells[0].text = str(row["week"])
        _cell_lines(cells[1], ([row["dates"]] if row["dates"] else []) + row["notes"], "" if has_dates else "[Dates]")
        cells[2].text = row["unit"] or ""
        _cell_lines(cells[3], [r.title for r in row["readings"]])
        _cell_lines(cells[4], [f"{a.title} (due {a.due_label})" if a.due_label else a.title for a in row["due"]])


def build_docx(ctx: dict, state: dict) -> bytes:
    doc = Document()
    section = doc.sections[0]
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(1))
    doc.styles["Normal"].font.size = Pt(11)
    doc.core_properties.title = ctx["title"]
    if ctx.get("instructor"):
        doc.core_properties.author = ctx["instructor"]

    doc.add_heading(ctx["title"], level=0)
    sub = " · ".join(x for x in (ctx.get("institution", ""), ctx.get("term_label", "")) if x)
    p = doc.add_paragraph()
    if sub:
        p.add_run(sub)
    else:
        _placeholder(p, "[Institution and term]")

    result = ctx.get("schedule")
    doc.add_heading("Course information", level=1)
    _field(doc, "Instructor", ctx.get("instructor", ""), "[Instructor name]")
    _field(doc, "Contact", placeholder="[Email and other ways to reach you]")
    _field(doc, "Office hours", placeholder="[Days, times, and place or link]")
    if result:
        when = ", ".join(result["meeting_days"])
        if result.get("start_time"):
            when += f", {result['start_time']}–{result['end_time']}"
        p = _field(doc, "Class meetings", when or "", "[Days and times]")
        p.add_run(" ")
        _placeholder(p, "[Room or link]")
        _field(doc, "Term dates", result["label"])
    else:
        _field(doc, "Class meetings", placeholder="[Days, times, and room or link]")

    doc.add_heading("Course description", level=1)
    _placeholder(doc.add_paragraph(), "[Course description]")

    doc.add_heading("Course materials", level=1)
    p = doc.add_paragraph("Every reading in this course is openly licensed or free to read online, so there is "
                          "nothing to buy. The readings are collected in the course reader: ")
    _placeholder(p, "[Link to the reader or where to find it]")

    sets = outcome_sets(state)
    if sets:
        doc.add_heading("Course outcomes", level=1)
        for o, texts in sets:
            if len(sets) > 1 or o.kind == "framework":
                doc.add_heading(" ".join(x for x in (o.institution, o.course, o.title) if x), level=2)
            _numbered(doc, texts)

    doc.add_heading("Schedule", level=1)
    if not result:
        _placeholder(doc.add_paragraph(), "[Set class dates in Aletheia to fill in the dates, or add them here.]")
    _schedule_table(doc, ctx)

    majors = [(u, a) for u in ctx["units"] for a in u.assignments if getattr(a, "role", "") in MAJOR_ROLES]
    if majors:
        doc.add_heading("Major assignments", level=1)
        for u, a in majors:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(a.title).bold = True
            p.add_run(f" ({u.title})")
            if getattr(a, "due_label", ""):
                p.add_run(f". Due {a.due_label}")
            p.add_run(". ")
            _placeholder(p, "[Weight or points]")
            if a.best_url:
                p.add_run(" ")
                _hyperlink(p, a.best_url, "Assignment")

    doc.add_heading("Course policies", level=1)
    for heading, text in POLICIES:
        doc.add_heading(heading, level=2)
        _placeholder(doc.add_paragraph(), text)

    doc.add_heading("Credits and licenses", level=1)
    doc.add_paragraph("The readings and assignments in this course come from these sources. "
                      "Each one keeps its own license.")
    items = list(ctx["readings"]) + list(ctx["assignments"])
    for item, p in zip(items, _numbered(doc, [i.credit for i in items])):
        if item.license_url:
            p.add_run(" ")
            _hyperlink(p, item.license_url, "License")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
