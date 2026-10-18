import io
import zipfile
from xml.etree import ElementTree as ET

import pytest
from django.test import Client

from catalog.models import Assignment, Pack, Reading


@pytest.fixture
def catalog(db):
    pack = Pack.objects.create(id="t", name="Test", version="0", config={})
    for rid, tier, verified in [("open", "adapt", True), ("nd", "verbatim", True), ("closed", "link-only", True),
                                ("unverified", "adapt", False)]:
        Reading.objects.create(id=rid, pack=pack, level="chapter", title=f"Reading {rid}", authors="Doe, Jane",
                               license="CC BY-NC-ND 4.0" if tier == "verbatim" else "CC BY 4.0",
                               license_family="CC BY-NC-ND" if tier == "verbatim" else "CC BY",
                               reuse_tier=tier, license_verified=verified, html_url=f"https://example.org/{rid}")
    Assignment.objects.create(id="essay", pack=pack, title="Research essay", kind="major assignment",
                              license="CC BY-NC 4.0", license_family="CC BY-NC", reuse_tier="adapt-nc",
                              license_verified=True, url="https://example.org/essay")
    return pack


STATE = {
    "title": "ENGL 101: Test & Course",
    "term": {"weeks": 10},
    "units": [
        {"key": "u1", "title": "Unit <One>", "start_week": 1, "end_week": 3, "goals": ["Read closely"],
         "readings": [{"id": "open", "selected": True}, {"id": "nd", "selected": True},
                      {"id": "closed", "selected": True}, {"id": "unverified", "selected": True},
                      {"id": "skipped", "selected": False}],
         "assignments": [{"id": "essay", "role": "major", "selected": True}]},
    ],
}


def test_reader_html_escapes_and_credits(catalog):
    resp = Client().post("/api/export/html", {"state": STATE}, content_type="application/json")
    assert resp.status_code == 200
    html = resp.content.decode()
    assert "Unit &lt;One&gt;" in html
    assert "Reading open" in html and "Licensed under CC BY 4.0" in html
    assert "linked, not copied" in html
    assert "NoDerivatives" in html


def test_unverified_license_is_treated_as_link_only(catalog):
    report = Client().post("/api/license-report", {"state": STATE}, content_type="application/json").json()
    assert report["tiers"]["link-only"] == 2
    assert report["unverified"] == ["unverified"]


def _check_package(z: zipfile.ZipFile, flavor: str) -> None:
    names = z.namelist()
    manifest = ET.fromstring(z.read("imsmanifest.xml"))
    ns = {"cc": "http://www.imsglobal.org/xsd/imsccv1p1/imscp_v1p1"}
    resources = {r.get("identifier") for r in manifest.iterfind(".//cc:resource", ns)}
    refs = [i.get("identifierref") for i in manifest.iterfind(".//cc:item", ns) if i.get("identifierref")]
    assert refs and all(r in resources for r in refs), "every module item must point at a resource"
    for r in manifest.iterfind(".//cc:resource", ns):
        for f in r.iterfind("cc:file", ns):
            assert f.get("href") in names
    if flavor == "canvas":
        assert "course_settings/canvas_export.txt" in names
        ET.fromstring(z.read("course_settings/module_meta.xml"))
        assert any(n.endswith("assignment_settings.xml") for n in names)
    else:
        assert not any(n.startswith("course_settings/") for n in names)


@pytest.mark.parametrize("flavor", ["canvas", "standard"])
def test_imscc_package_is_well_formed(catalog, flavor):
    resp = Client().post(f"/api/export/imscc?flavor={flavor}", {"state": STATE}, content_type="application/json")
    assert resp.status_code == 200
    _check_package(zipfile.ZipFile(io.BytesIO(resp.content)), flavor)


# ---------------------------------------------------------------- dated schedules

SCHEDULE = {"start_date": "2027-01-04", "meeting_days": ["Mon", "Wed"],
            "no_class": [{"date": "2027-01-11", "label": "Snow day"},
                         {"start": "2027-01-18", "end": "2027-01-22", "label": "Break"}],
            "timezone": "America/Los_Angeles"}
DATED = {**STATE, "title": "ENGL 101: Writing, Research; Inquiry", "term": {"weeks": 4, "fixed_start": 0, "fixed_end": 1},
         "schedule": SCHEDULE}


def post(path, state):
    return Client().post(path, {"state": state}, content_type="application/json")


def canvas_files(state):
    z = zipfile.ZipFile(io.BytesIO(post("/api/export/imscc?flavor=canvas", state).content))
    return z, {n: z.read(n).decode() for n in z.namelist()}


def test_exports_without_a_schedule_have_no_dates(catalog):
    html = post("/api/export/html", STATE).content.decode()
    assert '<span class="weeks">(Weeks 1–3)</span>' in html and "Due " not in html
    _, files = canvas_files(STATE)
    settings = next(v for k, v in files.items() if k.endswith("assignment_settings.xml"))
    assert "<due_at/>" in settings
    assert "<title>Weeks 1–3: Unit &lt;One&gt;</title>" in files["course_settings/module_meta.xml"]
    schedule_page = next(v for k, v in files.items() if k.endswith("course-schedule.html"))
    assert '<th scope="col">When</th>' in schedule_page


def test_reader_shows_dates_and_due_dates(catalog):
    html = post("/api/export/html", DATED).content.decode()
    # The break pushes the unit's third week to Jan 25–29.
    assert '<span class="weeks">(Weeks 1–3, Jan 4–29)</span>' in html
    assert '<p class="weeks">Weeks 1–3, Jan 4–29</p>' in html
    assert "Due Wed, Jan 27" in html


@pytest.mark.parametrize("flavor", ["canvas", "standard"])
def test_dated_imscc_is_well_formed(catalog, flavor):
    resp = post(f"/api/export/imscc?flavor={flavor}", DATED)
    _check_package(zipfile.ZipFile(io.BytesIO(resp.content)), flavor)


def test_dated_imscc_sets_due_dates_titles_and_schedule(catalog):
    z, files = canvas_files(DATED)
    settings = next(v for k, v in files.items() if k.endswith("assignment_settings.xml"))
    assert "<due_at>2027-01-27T23:59:00-08:00</due_at>" in settings
    ET.fromstring(settings)
    assert "<title>Weeks 1–3, Jan 4–29: Unit &lt;One&gt;</title>" in files["course_settings/module_meta.xml"]
    page = next(v for k, v in files.items() if k.endswith("course-schedule.html"))
    assert '<th scope="col">Dates</th>' in page
    assert "<td>Jan 4–8</td>" in page
    assert "No class Mon: Snow day" in page
    assert "<em>No class: Break</em>" in page
    assert "Research essay (Wed, Jan 27)" in page


def test_schedule_preview(catalog):
    assert post("/api/schedule", STATE).json() == {"ok": False, "warnings": [], "notes": []}
    r = post("/api/schedule", DATED).json()
    assert r["ok"] and r["weeks"][0]["label"] == "Jan 4–8"
    assert r["units"]["u1"]["due"] == {"essay": "2027-01-27"}
    assert r["breaks"][0]["label"] == "Break"


def unfold(text: str) -> list[str]:
    return text.replace("\r\n ", "").split("\r\n")


def events(lines: list[str]) -> list[dict]:
    out, cur = [], None
    for line in lines:
        if line == "BEGIN:VEVENT":
            cur = {}
        elif line == "END:VEVENT":
            out.append(cur)
            cur = None
        elif cur is not None:
            name, _, value = line.partition(":")
            cur[name] = value
    return out


def test_ics_needs_a_schedule(catalog):
    resp = post("/api/export/ics", STATE)
    assert resp.status_code == 400
    assert "start date" in resp.json()["detail"]


def test_ics_calendar(catalog):
    resp = post("/api/export/ics", DATED)
    assert resp.status_code == 200
    assert resp["Content-Type"].startswith("text/calendar")
    assert resp["Content-Disposition"].endswith('.ics"')
    raw = resp.content
    assert raw.endswith(b"\r\n") and b"\n" not in raw.replace(b"\r\n", b"")
    assert all(len(line) <= 75 for line in raw.split(b"\r\n"))
    lines = unfold(raw.decode())
    assert lines[0] == "BEGIN:VCALENDAR" and lines[-2] == "END:VCALENDAR"
    assert "VERSION:2.0" in lines
    evs = events(lines)
    meetings = [e for e in evs if "no class" not in e["SUMMARY"] and not e["SUMMARY"].startswith("Due")]
    # 4 weeks of Mon/Wed (the break week doesn't count), less the snow day: 7 meetings.
    assert len(meetings) == 7
    assert meetings[0]["DTSTART;VALUE=DATE"] == "20270104" and meetings[0]["DTEND;VALUE=DATE"] == "20270105"
    assert meetings[0]["SUMMARY"] == "ENGL 101: Writing\\, Research\\; Inquiry: Unit <One>"
    assert "Readings due:\\n- Reading open (https://example.org/open)" in meetings[0]["DESCRIPTION"]
    due = [e for e in evs if e["SUMMARY"].startswith("Due")]
    assert [(e["SUMMARY"], e["DTSTART;VALUE=DATE"]) for e in due] == [("Due: Research essay", "20270127")]
    off = sorted(e["SUMMARY"] for e in evs if "no class" in e["SUMMARY"])
    assert off == ["ENGL 101: Writing\\, Research\\; Inquiry: no class (Break)",
                   "ENGL 101: Writing\\, Research\\; Inquiry: no class (Snow day)"]
    uids = [e["UID"] for e in evs]
    assert len(set(uids)) == len(uids)
    # Re-exporting after a schedule change keeps the same UIDs, so calendars update in place.
    moved = {**DATED, "schedule": {**SCHEDULE, "no_class": [], "start_date": "2027-01-05"}}
    again = {e["UID"] for e in events(unfold(post("/api/export/ics", moved).content.decode()))}
    assert {e["UID"] for e in meetings} <= again
    assert due[0]["UID"] in again


def test_ics_timed_meetings_are_utc(catalog):
    state = {**DATED, "schedule": {**SCHEDULE, "start_time": "10:00", "end_time": "10:50"}}
    evs = events(unfold(post("/api/export/ics", state).content.decode()))
    first = next(e for e in evs if "DTSTART" in e)
    assert first["DTSTART"] == "20270104T180000Z" and first["DTEND"] == "20270104T185000Z"


def test_ics_folding_and_escaping():
    from builder.ics import escape_text, fold

    assert escape_text("a,b;c\\d\ne") == "a\\,b\\;c\\\\d\\ne"
    line = "SUMMARY:" + "é" * 80
    folded = fold(line)
    parts = folded.split("\r\n")
    assert all(len(p.encode()) <= 75 for p in parts)
    assert all(p.startswith(" ") for p in parts[1:])
    assert folded.replace("\r\n ", "") == line


def docx_text(resp):
    from docx import Document

    doc = Document(io.BytesIO(resp.content))
    return doc, [p.text for p in doc.paragraphs]


@pytest.fixture
def outcomes(catalog):
    from catalog.models import Outcome, OutcomeSet

    s = OutcomeSet.objects.create(id="test-101", pack=catalog, kind="course", institution="Test U", course="ENG 101",
                                  title="Writing")
    Outcome.objects.create(id="test-101:1", outcome_set=s, order=1, text="Write for real readers.")
    Outcome.objects.create(id="test-101:2", outcome_set=s, order=2, text="Revise.")
    return s


def test_docx_syllabus_with_schedule(outcomes):
    resp = post("/api/export/docx", {**DATED, "outcome_set_ids": ["test-101", "missing"]})
    assert resp.status_code == 200
    assert resp["Content-Disposition"].endswith('-syllabus.docx"')
    doc, paras = docx_text(resp)
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    for h in ["Course information", "Course description", "Course materials", "Course outcomes", "Schedule",
              "Major assignments", "Course policies", "Generative AI", "Title IX and basic needs", "Credits and licenses"]:
        assert h in headings
    assert doc.paragraphs[0].text == DATED["title"]
    assert "Instructor: [Instructor name]" in paras
    assert "Term dates: Jan 4 – Feb 5" in paras
    assert "Write for real readers." in paras and "Revise." in paras
    assert any("openly licensed or free to read" in p for p in paras)
    assert any(p.startswith("Research essay (Unit <One>). Due Wed, Jan 27.") for p in paras)
    assert any("Reading open" in p and "Licensed under CC BY 4.0" in p for p in paras)
    assert any("linked, not copied" in p for p in paras)
    table = doc.tables[0]
    rows = [[c.text for c in r.cells] for r in table.rows]
    assert rows[0] == ["Week", "Dates", "Unit", "Readings", "Assignments due"]
    assert rows[1][:3] == ["1", "Jan 4–8", "Unit <One>"]
    assert "No class Mon: Snow day" in rows[2][1]
    assert rows[3][2] == "No class: Break"
    assert any("Research essay (due Wed, Jan 27)" in r[4] for r in rows)
    readings_listed = "\n".join(r[3] for r in rows)
    for title in ["Reading open", "Reading nd", "Reading closed", "Reading unverified"]:
        assert title in readings_listed


def test_docx_syllabus_without_schedule_or_outcomes(catalog):
    resp = post("/api/export/docx", STATE)
    assert resp.status_code == 200
    doc, paras = docx_text(resp)
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "Course outcomes" not in headings
    assert "Schedule" in headings
    rows = [[c.text for c in r.cells] for r in doc.tables[0].rows]
    assert all(r[1] == "[Dates]" for r in rows[1:])
